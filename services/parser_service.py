"""
Сервис парсинга вложений КП (PDF/Excel)
"""
import os
from pathlib import Path
from typing import Dict, List, Optional
import pandas as pd
from loguru import logger


class ParserService:
    """Сервис парсинга PDF и Excel файлов"""
    
    def __init__(self, llm_service=None):
        self.llm = llm_service
    
    def parse_pdf(self, file_path: str) -> Dict:
        """
        Парсинг PDF файла
        
        Используется pdfplumber для извлечения таблиц
        """
        data = {
            "prices": [],
            "terms": [],
            "notes": ""
        }
        
        try:
            import pdfplumber
            
            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages):
                    # Извлечение таблиц
                    tables = page.extract_tables()
                    for table in tables:
                        parsed_table = self._parse_price_table(table)
                        data["prices"].extend(parsed_table)
                    
                    # Извлечение текста
                    text = page.extract_text()
                    if text:
                        data["notes"] += text + "\n"
            
            # Очистка текста
            data["notes"] = data["notes"][:2000]  # Ограничение
            
            logger.info(f"PDF парсинг завершён: {len(data['prices'])} цен найдено")
        
        except Exception as e:
            logger.error(f"Ошибка парсинга PDF: {e}")
            
            # Fallback на LLM
            if self.llm:
                logger.info("Попробую LLM fallback")
                data = self._parse_with_llm(file_path, "pdf")
        
        return data
    
    def parse_excel(self, file_path: str) -> Dict:
        """
        Парсинг Excel файла
        """
        data = {
            "prices": [],
            "terms": [],
            "notes": ""
        }
        
        try:
            # Чтение всех листов
            xls = pd.ExcelFile(file_path)
            
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name)
                
                # Попытка извлечь цены
                prices = self._parse_excel_table(df)
                data["prices"].extend(prices)
                
                # Текстовые данные
                data["terms"].append(f"Sheet: {sheet_name}")
            
            logger.info(f"Excel парсинг завершён: {len(data['prices'])} цен найдено")
        
        except Exception as e:
            logger.error(f"Ошибка парсинга Excel: {e}")
            
            # Fallback на LLM
            if self.llm:
                logger.info("Попробую LLM fallback")
                data = self._parse_with_llm(file_path, "excel")
        
        return data
    
    def parse_docx(self, file_path: str) -> Dict:
        """
        Парсинг Word документа
        """
        data = {
            "prices": [],
            "terms": [],
            "notes": ""
        }
        
        try:
            from docx import Document
            
            doc = Document(file_path)
            text = []
            
            for paragraph in doc.paragraphs:
                text.append(paragraph.text)
            
            data["notes"] = "\n".join(text)[:2000]
            
            # Извлечение таблиц
            for table in doc.tables:
                parsed = self._parse_word_table(table)
                data["prices"].extend(parsed)
            
            logger.info(f"DOCX парсинг завершён: {len(data['prices'])} цен найдено")
        
        except Exception as e:
            logger.error(f"Ошибка парсинга DOCX: {e}")
            
            if self.llm:
                data = self._parse_with_llm(file_path, "docx")
        
        return data
    
    def _parse_price_table(self, table: List[List]) -> List[Dict]:
        """Парсинг таблицы цен из PDF"""
        prices = []
        
        for row in table[1:]:  # Пропускаем заголовок
            if len(row) >= 2:
                service = row[0].strip() if row[0] else ""
                price = row[1].strip() if row[1] else ""
                
                if service and price:
                    prices.append({
                        "service": service,
                        "price": price,
                        "currency": "RUB"
                    })
        
        return prices
    
    def _parse_excel_table(self, df: pd.DataFrame) -> List[Dict]:
        """Парсинг таблицы из Excel"""
        prices = []
        
        # Поиск колонок с ценой и услугой
        price_col = None
        service_col = None
        
        for col in df.columns:
            col_lower = str(col).lower()
            if 'цена' in col_lower or 'price' in col_lower or 'cost' in col_lower:
                price_col = col
            if 'услуга' in col_lower or 'service' in col_lower or 'наименование' in col_lower:
                service_col = col
        
        if price_col and service_col:
            for _, row in df.iterrows():
                prices.append({
                    "service": str(row[service_col]),
                    "price": str(row[price_col]),
                    "currency": "RUB"
                })
        
        return prices
    
    def _parse_word_table(self, table) -> List[Dict]:
        """Парсинг таблицы из Word"""
        prices = []
        
        rows = table.rows
        for row in rows[1:]:  # Пропускаем заголовок
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) >= 2:
                prices.append({
                    "service": cells[0],
                    "price": cells[1],
                    "currency": "RUB"
                })
        
        return prices
    
    def _parse_with_llm(self, file_path: str, file_type: str) -> Dict:
        """
        Парсинг через LLM (fallback)
        """
        if not self.llm:
            return {"prices": [], "terms": [], "notes": "LLM not configured"}
        
        # Чтение содержимого файла
        try:
            if file_type == "pdf":
                import pdfplumber
                with pdfplumber.open(file_path) as pdf:
                    text = ""
                    for page in pdf.pages:
                        text += page.extract_text() or ""
            elif file_type == "excel":
                df = pd.read_excel(file_path)
                text = df.to_string()
            else:
                from docx import Document
                doc = Document(file_path)
                text = "\n".join([p.text for p in doc.paragraphs])
            
            # Запрос к LLM
            return self.llm.parse_kp_with_llm(text)
        
        except Exception as e:
            logger.error(f"Ошибка LLM парсинга: {e}")
            return {"prices": [], "terms": [], "notes": f"Error: {str(e)}"}
    
    def parse_file(self, file_path: str) -> Dict:
        """
        Автоматический выбор парсера по расширению
        """
        ext = Path(file_path).suffix.lower()
        
        if ext == ".pdf":
            return self.parse_pdf(file_path)
        elif ext in [".xlsx", ".xls"]:
            return self.parse_excel(file_path)
        elif ext == ".docx":
            return self.parse_docx(file_path)
        else:
            logger.warning(f"Неподдерживаемый формат: {ext}")
            return {"prices": [], "terms": [], "notes": ""}
